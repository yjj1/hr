# -*- coding:'utf8' -*-
#encoding=utf-8
from xml.dom import minidom

import docx
import zipfile
import types
import uuid
from orm import *
from connect_db import ConnectDb
from win32com import client as wc
import os
#global var
global connectDb

class ParseDoc:
    emailId = None
    def __init__(self):
        print 'init'

    def setEmailId(self, eId):
        self.emailId = eId

    def parse(self,file_path):
        path_array = os.path.splitext(file_path)
        file_name = path_array[0]
        suffix = path_array[1]
        sDoc = '.doc'
        if suffix == sDoc:
            file_path = self.doc2docx(file_path)
        fullText = []
        doc = docx.Document(file_path)
        tables = doc.tables

        t = tables[0]
        for i in range(0, 330):
            content = t.cell(0, i).text
            print content, '(',0,',',i, ')'

        sUuid = str(uuid.uuid4())
        data = ResumeBaseInfo(sUuid)

        name = t.cell(0,1).text
        sex = t.cell(0,5).text
        birthday = t.cell(0,11).text
        eduBack = t.cell(0,16).text
        height = t.cell(0, 20).text
        health = t.cell(0, 26).text
        politicsStatus = t.cell(0, 31).text
        nativePlace = t.cell(0,35).text
        marriage = t.cell(0,41).text
        pro = t.cell(0,54).text
        school = t.cell(0,46).text
        address = t.cell(0,61).text
        idCard = t.cell(0,76).text
        phone = t.cell(0,93).text
        homePhone = t.cell(0,99).text
        czSchool = t.cell(0,123).text
        czPro = t.cell(0,130).text
        czDate = t.cell(0,132).text

        gzSchool = t.cell(0,138).text
        gzPro = t.cell(0,145).text
        gzDate = t.cell(0,147).text

        dzSchool = t.cell(0,153).text
        dzPro = t.cell(0,160).text
        dzDate = t.cell(0,162).text

        bkSchool = t.cell(0,168).text
        bkPro = t.cell(0,175).text
        bkDate = t.cell(0,177).text

        yjsSchool = t.cell(0, 183).text
        yjsPro = t.cell(0, 190).text
        yjsDate = t.cell(0, 192).text

        bigBang = t.cell(0,196).text

        job = t.cell(0,287).text
        isAddressControl = t.cell(0,308).text
        hobbies = t.cell(0, 271).text
        #家庭
        f1 = Family(str(uuid.uuid4()))
        f1.resumeId = data.resumeId
        f1.name=t.cell(0,226).text
        f1.relationShip=t.cell(0,229).text
        f1.politicsStatus=t.cell(0,231).text
        f1.company=t.cell(0,235).text
        f1.job=t.cell(0,239).text

        f2 = Family(str(uuid.uuid4()))
        f2.resumeId = data.resumeId
        f2.name = t.cell(0, 241).text
        f2.relationShip = t.cell(0, 244).text
        f2.politicsStatus = t.cell(0, 246).text
        f2.company = t.cell(0, 250).text
        f2.job = t.cell(0, 254).text

        f3 = Family(str(uuid.uuid4()))
        f3.resumeId = data.resumeId
        f3.name = t.cell(0, 256).text
        f3.relationShip = t.cell(0, 259).text
        f3.politicsStatus = t.cell(0, 261).text
        f3.company = t.cell(0, 265).text
        f3.job = t.cell(0, 269).text

        # f4 = Family(str(uuid.uuid4()))
        # f4.resumeId = data.resumeId
        # f4.name = t.cell(0, 273).text
        # f4.relationShip = t.cell(0, 275).text
        # f4.politicsStatus = t.cell(0, 278).text
        # f4.company = t.cell(0, 281).text
        # f4.job = t.cell(0, 287).text

        data.name = name
        data.sex = sex
        data.birthday = birthday
        data.eduBack = eduBack
        data.phone = phone
        data.homePhone = homePhone
        data.resumePath = file_path
        data.politicsStatus = politicsStatus
        data.nativePlace = nativePlace
        data.address = address
        data.bigBang = bigBang
        data.height = height
        data.healthy = health
        data.marriage = marriage
        data.idCard = idCard
        data.job = job
        data.isAddressControl = isAddressControl
        data.hobbies = hobbies
        data.eduDegree = pro
        data.school = school

        if self.emailId is not None:
            data.emailId = self.emailId

        eduDataCz = Edu(str(uuid.uuid4()))
        eduDataCz.resumeId = data.resumeId
        eduDataCz.eduDate = czDate
        eduDataCz.eduPro = czPro
        eduDataCz.eduType = '1'
        eduDataCz.schoolName = czSchool

        eduDataGz = Edu(str(uuid.uuid4()))
        eduDataGz.resumeId = data.resumeId
        eduDataGz.eduDate = gzDate
        eduDataGz.eduPro = gzPro
        eduDataGz.eduType = '2'
        eduDataGz.schoolName = gzSchool

        eduDataDz = Edu(str(uuid.uuid4()))
        eduDataDz.resumeId = data.resumeId
        eduDataDz.eduDate = dzDate
        eduDataDz.eduPro = dzPro
        eduDataDz.eduType = '3'
        eduDataDz.schoolName = dzSchool

        eduDataBk = Edu(str(uuid.uuid4()))
        eduDataBk.resumeId = data.resumeId
        eduDataBk.eduDate = bkDate
        eduDataBk.eduPro = bkPro
        eduDataBk.eduType = '4'
        eduDataBk.schoolName = bkSchool

        eduDataYjs = Edu(str(uuid.uuid4()))
        eduDataYjs.resumeId = data.resumeId
        eduDataYjs.eduDate = yjsDate
        eduDataYjs.eduPro = yjsPro
        eduDataYjs.eduType = '5'
        eduDataYjs.schoolName = yjsSchool

        connectDb = ConnectDb()
        session = connectDb.db_session
        session.add(data)
        session.add(f1)
        session.add(f2)
        session.add(f3)
        # session.add(f4)
        session.add(eduDataCz)
        session.add(eduDataGz)
        session.add(eduDataDz)
        session.add(eduDataBk)
        session.add(eduDataYjs)

        session.commit()

        print t.cell(0,8).text
        print t.cell(0,9).text
        print t.cell(0,10).text
        print t.cell(0,11).text

        #如果从邮件中来，则更新邮件状态
        if self.emailId is not None:
            session.query(OriginEmail).filter(OriginEmail.emailId == self.emailId).update({
                OriginEmail.status:'1'
            })
            session.commit()

        return fullText

    def zipParse(self, file_path):
        document = zipfile.ZipFile(file_path)
        xml_content = document.read('word/document.xml')
        reparsed = minidom.parseString(xml_content)
        print reparsed.toprettyxml(indent="   ", encoding="utf-8")

    def doc2docx(self, file_path):
        new_file_path = file_path + 'x'

        word = wc.Dispatch('Word.Application')
        doc = word.Documents.Open(file_path)
        doc.SaveAs(new_file_path, FileFormat=16)
        doc.Close()

        word.Quit()

        return new_file_path

if __name__ == '__main__':
    # text = ParseDoc().parse('F:\\apache-tomcat-7.0.72\\bin\\hr_file\\0d1731b7-05ea-4257-b96e-02eeccea1435玉环农商行2018_叶坚静.doc')
    text = ParseDoc().parse('D:\\test2.doc')
# ParseDoc().zipParse('/users/jianjingye/test2.doc')
    # for t in text:
    #     print t.text